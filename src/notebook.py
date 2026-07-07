from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document

from src.models import LookupResult


class NotebookError(RuntimeError):
    """Raised when a Word notebook cannot be created or updated."""


def normalize_saved_query(text: str) -> str:
    return " ".join(text.strip().casefold().split())


def validate_notebook_path(path: Path) -> None:
    if path.suffix.lower() != ".docx":
        raise NotebookError("单词本文件必须是 .docx 格式")
    if not path.parent.exists():
        raise NotebookError(f"单词本目录不存在：{path.parent}")


class WordNotebook:
    def __init__(self, path: Path) -> None:
        self.path = path

    def append(self, result: LookupResult) -> bool:
        validate_notebook_path(self.path)
        try:
            document = Document(self.path) if self.path.exists() else self._new_document()
            table = self._get_or_create_table(document)
            if self._contains_query(table, result.query):
                return False
            row = table.add_row().cells
            row[0].text = result.query
            row[1].text = result.translation
            row[2].text = result.phonetic_us
            row[3].text = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            document.save(self.path)
            return True
        except PermissionError as exc:
            raise NotebookError("无法写入单词本，文件可能正在被 Word 打开或权限不足") from exc
        except OSError as exc:
            raise NotebookError(f"写入单词本失败：{self.path}") from exc

    def _new_document(self) -> Document:
        document = Document()
        document.add_heading("离线选词单词本", level=1)
        return document

    def _get_or_create_table(self, document: Document):
        if document.tables:
            table = document.tables[0]
            if len(table.columns) >= 4:
                return table

        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        headers = table.rows[0].cells
        headers[0].text = "原文"
        headers[1].text = "中文释义"
        headers[2].text = "美式音标"
        headers[3].text = "保存时间"
        return table

    def _contains_query(self, table, query: str) -> bool:
        normalized_query = normalize_saved_query(query)
        for row in table.rows[1:]:
            if row.cells and normalize_saved_query(row.cells[0].text) == normalized_query:
                return True
        return False
