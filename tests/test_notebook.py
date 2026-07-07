from __future__ import annotations

from pathlib import Path

from docx import Document

from src.models import LookupResult
from src.notebook import WordNotebook


def test_word_notebook_appends_lookup_result(tmp_path: Path) -> None:
    notebook_path = tmp_path / "words.docx"
    result = LookupResult(
        query="hello",
        translation="int. 你好",
        phonetic_us="həˈloʊ",
        found=True,
    )

    WordNotebook(notebook_path).append(result)

    document = Document(notebook_path)
    table = document.tables[0]
    cells = table.rows[1].cells
    assert cells[0].text == "hello"
    assert cells[1].text == "int. 你好"
    assert cells[2].text == "həˈloʊ"


def test_word_notebook_skips_duplicate_query(tmp_path: Path) -> None:
    notebook_path = tmp_path / "words.docx"
    notebook = WordNotebook(notebook_path)

    first_saved = notebook.append(
        LookupResult(
            query="hello",
            translation="int. 你好",
            phonetic_us="həˈloʊ",
            found=True,
        )
    )
    second_saved = notebook.append(
        LookupResult(
            query=" Hello ",
            translation="新的释义",
            phonetic_us="həˈloʊ",
            found=True,
        )
    )

    document = Document(notebook_path)
    table = document.tables[0]
    assert first_saved is True
    assert second_saved is False
    assert len(table.rows) == 2
